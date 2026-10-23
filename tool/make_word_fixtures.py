"""Faz 22 Word örnek seti (python-docx 1.x; DUZENLEME_PLANI §8 madde 2).

    python3 tool/make_word_fixtures.py

-> test/fixtures/office/stiller_basliklar.docx, liste_duzeyli.docx,
   arapca_rtl.docx. python-docx'in varsayılan şablonu Word'ün tam
   styles.xml / numbering.xml / theme1.xml'ini taşır: gerçek belge gibi.
"""
import io
import os
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

OUT = os.path.join(os.path.dirname(__file__), '..', 'test', 'fixtures', 'office')


def png(w=40, h=24):
    raw = b''.join(b'\x00' + bytes([200, 60, 60] * w) for _ in range(h))
    def chunk(t, d):
        c = struct.pack('>I', len(d)) + t + d
        return c + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def stiller():
    d = Document()
    d.add_heading('Proje Raporu 2026', 0)
    d.add_heading('Giriş', level=1)
    p = d.add_paragraph('Bu belge ')
    p.add_run('kalın').bold = True
    p.add_run(', ')
    p.add_run('italik').italic = True
    r = p.add_run(' ve kırmızı')
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.add_run(' metin içerir. Türkçe: ğüşıöçİ.')
    d.add_heading('Ayrıntılar', level=2)
    c = d.add_paragraph('Ortalanmış paragraf.')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph('Birinci madde', style='List Bullet')
    d.add_paragraph('İkinci madde', style='List Bullet')
    d.add_paragraph('Adım bir', style='List Number')
    d.add_paragraph('Adım iki', style='List Number')
    t = d.add_table(rows=2, cols=2)
    t.style = 'Table Grid'
    for i, txt in enumerate(['Ad', 'Şehir', 'Ayşe', 'İzmir']):
        t.cell(i // 2, i % 2).text = txt
    d.add_paragraph('Resim aşağıda:')
    d.add_picture(io.BytesIO(png()), width=Inches(1.5))
    s = d.add_paragraph('Sekmeli\tsatır ve ')
    s.add_run().add_break()
    s.add_run('satır sonu.')
    d.add_paragraph('Son paragraf.')
    d.save(os.path.join(OUT, 'stiller_basliklar.docx'))


def liste():
    d = Document()
    d.add_heading('Liste düzeyleri', 1)
    for txt, style in [('Meyveler', 'List Bullet'), ('Elma', 'List Bullet 2'),
                       ('Armut', 'List Bullet 2'), ('Sebzeler', 'List Bullet'),
                       ('Bir', 'List Number'), ('İki', 'List Number'),
                       ('İki-a', 'List Number 2'), ('Üç', 'List Number')]:
        d.add_paragraph(txt, style=style)
    d.add_paragraph('Liste sonrası düz paragraf.')
    d.save(os.path.join(OUT, 'liste_duzeyli.docx'))


def arapca():
    d = Document()
    for txt in ['تقرير المشروع', 'هذا نص عربي من اليمين إلى اليسار.', 'سطر ثالث مع رقم 2026.']:
        p = d.add_paragraph(txt)
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
        for r in p.runs:
            rtl = OxmlElement('w:rtl')
            r._r.get_or_add_rPr().append(rtl)
    d.add_paragraph('Mixed: English satırı.')
    d.save(os.path.join(OUT, 'arapca_rtl.docx'))


if __name__ == '__main__':
    stiller()
    liste()
    arapca()
    print('ok')
